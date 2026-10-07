import io
from datetime import datetime

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

_FOOTER = "MuzMugalim Bot  |  @muzmugalim"
_W, _H = 960, 540          # 16:9, in points
_MARGIN = 40
_IMG = 360                 # square picture size


def _place(shape, left, top, width, height):
    shape.left, shape.top, shape.width, shape.height = (
        Pt(left), Pt(top), Pt(width), Pt(height)
    )


def build_pptx(title: str, section_label: str, slides: list, images: dict = None) -> tuple:
    """Build a .pptx from [{"title": str, "bullets": [str]}].

    images: optional {slide_index: png_bytes}; -1 is the cover picture, 0.. are content slides.
    Returns (bytes_io, filename).
    """
    images = images or {}
    prs = Presentation()
    prs.slide_width = Pt(_W)
    prs.slide_height = Pt(_H)

    # Cover
    cover = prs.slides.add_slide(prs.slide_layouts[0])
    has_cover_img = -1 in images
    text_w = (_W - 2 * _MARGIN - _IMG - 30) if has_cover_img else (_W - 2 * _MARGIN)
    _place(cover.shapes.title, _MARGIN, 140, text_w, 170)
    cover.shapes.title.text_frame.paragraphs[0].font.size = Pt(40)
    sub = cover.placeholders[1]
    _place(sub, _MARGIN, 330, text_w, 100)
    sub.text = f"{section_label}  •  {datetime.now().strftime('%d.%m.%Y')}\n{_FOOTER}"
    for para in sub.text_frame.paragraphs:
        para.font.size = Pt(18)
        para.font.color.rgb = RGBColor(90, 90, 100)
    if has_cover_img:
        cover.shapes.add_picture(
            io.BytesIO(images[-1]), Pt(_W - _MARGIN - _IMG), Pt((_H - _IMG) / 2), Pt(_IMG), Pt(_IMG)
        )

    # Content slides
    for idx, slide in enumerate(slides):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        has_img = idx in images
        body_w = (_W - 2 * _MARGIN - _IMG - 30) if has_img else (_W - 2 * _MARGIN)
        _place(s.shapes.title, _MARGIN, 24, _W - 2 * _MARGIN, 90)
        s.shapes.title.text = slide["title"]
        s.shapes.title.text_frame.paragraphs[0].font.size = Pt(34)
        body = s.placeholders[1]
        _place(body, _MARGIN, 130, body_w, _H - 130 - 40)
        tf = body.text_frame
        tf.clear()
        tf.word_wrap = True
        for i, bullet in enumerate(slide["bullets"]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = bullet
            para.font.size = Pt(24 if has_img else 26)
        if has_img:
            s.shapes.add_picture(
                io.BytesIO(images[idx]), Pt(_W - _MARGIN - _IMG), Pt(130), Pt(_IMG), Pt(_IMG)
            )

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    safe = "".join(c for c in title if c.isalnum() or c in " _-")[:40].strip() or "presentation"
    return buf, f"{safe}.pptx"
