import os
import sqlite3
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

_tmp = tempfile.NamedTemporaryFile(suffix=".db", delete=False)
_tmp.close()
os.environ["DB_FILE"] = _tmp.name
os.environ["ADMIN_CHAT_ID"] = "999"
os.environ.setdefault("GEMINI_API_KEY", "test-key")

import handlers  # noqa: E402
import storage  # noqa: E402
from config import MType, SECTIONS  # noqa: E402

SLIDES = [{"title": "Кіріспе", "bullets": ["а", "б"]}, {"title": "Қорытынды", "bullets": ["в"]}]


def _output_of(name):
    return next(m.output for m in SECTIONS["mektep"]["materials"] if m.name == name)


def _update(uid, text):
    msg = MagicMock()
    msg.text = text
    sent = MagicMock()
    sent.delete = AsyncMock()
    msg.reply_text = AsyncMock(return_value=sent)
    msg.reply_document = AsyncMock()
    msg.reply_photo = AsyncMock()
    user = SimpleNamespace(id=uid, username="u", first_name="T", last_name=None)
    return SimpleNamespace(effective_user=user, message=msg)


def _ctx(name, output):
    return SimpleNamespace(user_data={
        "waiting_topic": True, "section": "mektep", "material_name": name,
        "material_name_ru": "", "material_type": MType.TEXT,
        "material_output": output, "page": 0, "lang": "kz",
    })


class FlowTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        with sqlite3.connect(storage.DB_FILE) as c:
            c.execute("DELETE FROM users")

    def _user(self, uid, plan):
        storage.ensure_user(SimpleNamespace(id=uid, username="u", first_name="T", last_name=None))
        storage.activate_plan(uid, plan)

    def test_every_tariff_card_says_paid(self):
        for plan in ("basic", "standard", "premium"):
            self.assertIn("АҚЫЛЫ", handlers.TARIF_CARDS[plan])
        self.assertIn("ақылы", handlers.TARIF_INTRO)

    def test_no_plan_is_not_labelled_free(self):
        icon, name = handlers._plan_label("free")
        self.assertNotIn("Тегін", name)
        self.assertIn("Тариф жоқ", name)
        self.assertEqual(handlers._plan_label("premium")[1], "Премиум (ақылы)")

    def test_materials_have_outputs(self):
        self.assertEqual(_output_of("Презентация"), "pptx")
        self.assertEqual(_output_of("Көрнекілік"), "visual")
        self.assertEqual(_output_of("Ноталар"), "notes")

    async def test_presentation_basic_is_text_only_with_upgrade_hint(self):
        self._user(1, "basic")
        upd, ctx = _update(1, "Домбыра"), _ctx("Презентация", "pptx")
        gen_images = AsyncMock(return_value={})
        with patch.object(handlers.generators, "gen_slides", AsyncMock(return_value=SLIDES)),              patch.object(handlers.generators, "gen_slide_images", gen_images):
            await handlers.text_message_handler(upd, ctx)
        upd.message.reply_document.assert_awaited_once()
        self.assertTrue(upd.message.reply_document.await_args.kwargs["filename"].endswith(".pptx"))
        gen_images.assert_not_awaited()
        sent = " ".join(str(c.args[0]) for c in upd.message.reply_text.await_args_list if c.args)
        self.assertIn("Стандарт", sent)

    async def test_presentation_standard_gets_pictures_and_uses_one_poster_unit(self):
        import io as _io
        from PIL import Image
        from pptx import Presentation
        png = _io.BytesIO()
        Image.new("RGB", (32, 32), (10, 120, 200)).save(png, "PNG")
        self._user(5, "standard")
        upd, ctx = _update(5, "Домбыра"), _ctx("Презентация", "pptx")
        with patch.object(handlers.generators, "gen_slides", AsyncMock(return_value=SLIDES)),              patch.object(handlers.generators, "gen_slide_images",
                          AsyncMock(return_value={-1: png.getvalue(), 0: png.getvalue()})):
            await handlers.text_message_handler(upd, ctx)
        doc = upd.message.reply_document.await_args.kwargs["document"]
        prs = Presentation(_io.BytesIO(doc.getvalue()))
        pics = [sh for s in prs.slides for sh in s.shapes if sh.shape_type == 13]
        self.assertEqual(len(pics), 2)
        self.assertEqual(storage.get_remaining(5, "poster")[0], 1)

    async def test_notes_send_sheet_music_photo_and_text(self):
        import generators as g
        self._user(6, "basic")
        melody = g.parse_melody(
            '{"title":"Тест","time":"4/4","measures":[[{"p":"C4","d":1,"l":"а"},'
            '{"p":"D4","d":1,"l":"б"},{"p":"E4","d":2,"l":"в"}]]}'
        )
        upd, ctx = _update(6, "Көктем"), _ctx("Ноталар", "notes")
        with patch.object(handlers.generators, "gen_text", AsyncMock(return_value="сипаттама")),              patch.object(handlers.generators, "gen_melody", AsyncMock(return_value=melody)):
            await handlers.text_message_handler(upd, ctx)
        upd.message.reply_photo.assert_awaited_once()
        self.assertEqual(storage.get_remaining(6, "poster")[0], 0)  # notation costs no picture quota

    async def test_notes_text_still_sent_when_melody_fails(self):
        self._user(7, "basic")
        upd, ctx = _update(7, "Көктем"), _ctx("Ноталар", "notes")
        with patch.object(handlers.generators, "gen_text", AsyncMock(return_value="сипаттама")),              patch.object(handlers.generators, "gen_melody", AsyncMock(side_effect=ValueError("bad"))):
            await handlers.text_message_handler(upd, ctx)
        upd.message.reply_photo.assert_not_awaited()
        sent = " ".join(str(c.args[0]) for c in upd.message.reply_text.await_args_list if c.args)
        self.assertIn("сипаттама", sent)

    async def test_visual_standard_gets_text_and_image_and_uses_poster_quota(self):
        self._user(2, "standard")
        upd, ctx = _update(2, "Домбыра"), _ctx("Көрнекілік", "visual")
        with patch.object(handlers.generators, "gen_text", AsyncMock(return_value="мәтін")), \
             patch.object(handlers.generators, "gen_poster", AsyncMock(return_value=b"\x89PNG")):
            await handlers.text_message_handler(upd, ctx)
        upd.message.reply_photo.assert_awaited_once()
        self.assertEqual(storage.get_remaining(2, "poster")[0], 1)

    async def test_visual_basic_gets_text_and_upgrade_hint_no_image(self):
        self._user(3, "basic")
        upd, ctx = _update(3, "Домбыра"), _ctx("Көрнекілік", "visual")
        gen_poster = AsyncMock(return_value=b"x")
        with patch.object(handlers.generators, "gen_text", AsyncMock(return_value="мәтін")), \
             patch.object(handlers.generators, "gen_poster", gen_poster):
            await handlers.text_message_handler(upd, ctx)
        gen_poster.assert_not_awaited()
        upd.message.reply_photo.assert_not_awaited()
        sent = " ".join(str(c.args[0]) for c in upd.message.reply_text.await_args_list if c.args)
        self.assertIn("Стандарт", sent)

    async def test_visual_image_failure_does_not_break_text(self):
        self._user(4, "standard")
        upd, ctx = _update(4, "Домбыра"), _ctx("Көрнекілік", "visual")
        with patch.object(handlers.generators, "gen_text", AsyncMock(return_value="мәтін")), \
             patch.object(handlers.generators, "gen_poster", AsyncMock(side_effect=RuntimeError("boom"))):
            await handlers.text_message_handler(upd, ctx)
        upd.message.reply_photo.assert_not_awaited()
        self.assertEqual(storage.get_remaining(4, "poster")[0], 0)  # failed image not charged
        self.assertEqual(storage.get_remaining(4, "text")[0], 1)


if __name__ == "__main__":
    unittest.main()
