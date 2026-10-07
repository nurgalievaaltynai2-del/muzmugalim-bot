import io
import os
import unittest

os.environ.setdefault("GEMINI_API_KEY", "test-key")

from pptx import Presentation

import generators
import pptx_gen


SLIDES = [
    {"title": "Домбыра", "bullets": ["Ұлттық аспап", "Екі ішекті"]},
    {"title": "Тарихы", "bullets": ["Көне заман", "Күй өнері", "Ахмет Жұбанов"]},
]


class PptxTests(unittest.TestCase):
    def test_build_pptx_has_title_plus_content_slides(self):
        buf, fname = pptx_gen.build_pptx("Домбыра", "Мектеп", SLIDES)
        self.assertTrue(fname.endswith(".pptx"))
        prs = Presentation(io.BytesIO(buf.getvalue()))
        self.assertEqual(len(prs.slides), 1 + len(SLIDES))
        texts = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame]
        self.assertTrue(any("Ахмет Жұбанов" in t for t in texts))

    def test_parse_slides_accepts_fenced_json(self):
        raw = '```json\n{"slides": [{"title": "A", "bullets": ["x", "y"]}]}\n```'
        self.assertEqual(
            generators.parse_slides(raw),
            [{"title": "A", "bullets": ["x", "y"]}],
        )

    def test_parse_slides_rejects_garbage(self):
        with self.assertRaises(ValueError):
            generators.parse_slides("not json")
        with self.assertRaises(ValueError):
            generators.parse_slides('{"slides": []}')


if __name__ == "__main__":
    unittest.main()
